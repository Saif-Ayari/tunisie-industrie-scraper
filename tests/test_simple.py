from __future__ import annotations

import json
from pathlib import Path
import shutil
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from crm.simple import (
    CRM_MANAGED_HEADERS,
    SOURCE_ID,
    SimpleExportError,
    export_simple_records,
    load_simple_template_schema,
    transform_records,
    verify_simple_workbook,
)
from models import CompanyRecord


ROOT = Path(__file__).resolve().parents[1]
SCHEMA_FIXTURE = ROOT / "tests" / "fixtures" / "simple_company_schema.json"


def schema_headers() -> tuple[str, ...]:
    return tuple(json.loads(SCHEMA_FIXTURE.read_text(encoding="utf-8"))["headers"])


def make_template(path: Path) -> None:
    payload = json.loads(SCHEMA_FIXTURE.read_text(encoding="utf-8"))
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = payload["sheet"]
    worksheet.append(payload["headers"])
    workbook.save(path)


def sample_record(**overrides: object) -> CompanyRecord:
    values: dict[str, object] = {
        "short_name": "7 M",
        "company_name": "Société Étoile",
        "manager": "Élodie Ben Saïd",
        "activities": "Milk, cheese and yogurt manufacturing.",
        "products": "Lait et yaourt",
        "factory_address": "DISTRICT INDUSTRIEL - 4030 - ENFIDHA",
        "governorate": "Sousse",
        "delegation": "Enfidha",
        "phone": "(216) - 70 578 142 / 74 493 691",
        "fax": "(216) - 71 000 111",
        "email": "primary@example.tn",
        "url": "http://example.tn",
        "market": "Off-shore",
        "foreign_participant_country": "France",
        "created": "5/20/2019",
        "share_capital_dt": "1\xa0500\xa0000",
        "employees": "18",
        "source_id": "00017",
        "source_url": "https://www.tunisieindustrie.nat.tn/en/dbi.asp?action=result&ident=17",
        "scraped_at": "2026-10-01T13:00:00Z",
    }
    values.update(overrides)
    return CompanyRecord(**values)


class SimpleExportTests(unittest.TestCase):
    def test_exact_template_headers_and_all_mappings(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            template = directory_path / "template.xlsx"
            output = directory_path / "simple.xlsx"
            make_template(template)
            schema = load_simple_template_schema(template)
            self.assertEqual(schema.headers, schema_headers())

            report = export_simple_records([sample_record()], template, output)
            summary = verify_simple_workbook(output, schema, 1)
            self.assertEqual(report.rows_transformed, 1)
            self.assertEqual(summary["columns"], 40)

            workbook = load_workbook(output, data_only=False)
            worksheet = workbook["Export"]
            values = {header: worksheet.cell(2, index).value for index, header in enumerate(schema.headers, 1)}
            self.assertEqual(values["Nom"], "Société Étoile")
            self.assertEqual(values["Nom court"], "7 M")
            self.assertEqual(values["Responsable"], "Élodie Ben Saïd")
            self.assertEqual(values["Activités"], "Milk, cheese and yogurt manufacturing.")
            self.assertEqual(values["Produits"], "Lait et yaourt")
            self.assertEqual(values["Adresse"], "DISTRICT INDUSTRIEL - 4030 - ENFIDHA")
            self.assertEqual(values["Ville"], "Enfidha")
            self.assertEqual(values["État"], "Sousse")
            self.assertEqual(values["Code postal"], "4030")
            self.assertIsNone(values["Pays"])
            self.assertEqual(values["Numéro de téléphone"], "70 578 142")
            self.assertEqual(values["Code"], "TN")
            self.assertEqual(values["Indicatif"], "+216")
            self.assertEqual(values["Téléphone"], "+21674 493 691 (TN)")
            self.assertEqual(values["Fax"], "(216) - 71 000 111")
            self.assertEqual(values["Email principal"], "primary@example.tn")
            self.assertIsNone(values["Courriels"])
            self.assertEqual(values["Nom de domaine / URL du lien"], "http://example.tn")
            self.assertIsNone(values["Nom de domaine / Libellé du lien"])
            self.assertEqual(values["Marché"], "Off-shore")
            self.assertEqual(values["Pays participant étranger"], "France")
            self.assertEqual(values["Date de création entreprise"], "5/20/2019")
            self.assertEqual(values["Capital social DT"], "1\xa0500\xa0000")
            self.assertEqual(values["Effectif"], 18)
            self.assertEqual(values[SOURCE_ID], "00017")
            self.assertEqual(values["Source Tunisie Industrie / Libellé du lien"], "Tunisie Industrie")
            self.assertEqual(values["Source Tunisie Industrie / URL du lien"], sample_record().source_url)
            self.assertIsNone(values["Source Tunisie Industrie / Lien"])
            self.assertIsNone(values["Date de création"])
            self.assertIsNone(values["ID de l'enregistrement"])
            self.assertIsNone(values["Créé par"])
            self.assertIsNone(values["Propriétaire du compte"])
            self.assertNotIn("Scraped At", schema.headers)
            workbook.close()

    def test_international_phone_and_multiple_emails_are_preserved(self) -> None:
        schema = load_simple_template_schema(self._template())
        rows, report = transform_records(
            [
                sample_record(
                    phone="+33 1 42 68 53 00 / +33 1 43 11 22 33",
                    email="a@example.tn; b@example.tn; c@example.tn",
                )
            ],
            schema,
        )
        self.assertFalse(report.warnings)
        self.assertEqual(rows[0]["Numéro de téléphone"], "+33 1 42 68 53 00")
        self.assertIsNone(rows[0]["Code"])
        self.assertIsNone(rows[0]["Indicatif"])
        self.assertEqual(rows[0]["Téléphone"], "+33 1 43 11 22 33")
        self.assertEqual(rows[0]["Email principal"], "a@example.tn")
        self.assertEqual(rows[0]["Courriels"], "b@example.tn\nc@example.tn")

    def test_tunisian_additional_phones_use_twenty_newline_contract(self) -> None:
        schema = load_simple_template_schema(self._template())
        rows, report = transform_records(
            [sample_record(phone="70 578 142 / 74 493 691 / 78 469 362")],
            schema,
        )
        self.assertFalse(report.warnings)
        self.assertEqual(rows[0]["Numéro de téléphone"], "70 578 142")
        self.assertEqual(rows[0]["Code"], "TN")
        self.assertEqual(rows[0]["Indicatif"], "+216")
        self.assertEqual(
            rows[0]["Téléphone"],
            "+21674 493 691 (TN)\n+21678 469 362 (TN)",
        )

    def test_malformed_phone_is_reported_without_losing_valid_numbers(self) -> None:
        schema = load_simple_template_schema(self._template())
        rows, report = transform_records(
            [sample_record(phone="70 578 142 / not-a-phone / 74 493 691")],
            schema,
        )
        self.assertTrue(any("malformed phone" in warning for warning in report.warnings))
        self.assertEqual(rows[0]["Numéro de téléphone"], "70 578 142")
        self.assertEqual(rows[0]["Téléphone"], "+21674 493 691 (TN)")

    def test_blank_phone_stays_blank(self) -> None:
        schema = load_simple_template_schema(self._template())
        rows, report = transform_records([sample_record(phone=None)], schema)
        self.assertFalse(report.warnings)
        self.assertIsNone(rows[0]["Numéro de téléphone"])
        self.assertIsNone(rows[0]["Code"])
        self.assertIsNone(rows[0]["Indicatif"])
        self.assertIsNone(rows[0]["Téléphone"])

    def test_malformed_employees_warn_and_remain_blank(self) -> None:
        schema = load_simple_template_schema(self._template())
        rows, report = transform_records([sample_record(employees="about 18")], schema)
        self.assertIsNone(rows[0]["Effectif"])
        self.assertTrue(any("malformed Employees" in warning for warning in report.warnings))

    def test_duplicate_source_ids_are_rejected(self) -> None:
        schema = load_simple_template_schema(self._template())
        with self.assertRaises(SimpleExportError):
            transform_records([sample_record(), sample_record(company_name="Other")], schema)

    def test_deterministic_source_id_order_and_blank_values(self) -> None:
        schema = load_simple_template_schema(self._template())
        records = [
            sample_record(source_id="29", company_name="B"),
            sample_record(source_id="17", company_name="A", email=None, url=None, products=None),
        ]
        first, first_report = transform_records(records, schema)
        second, second_report = transform_records(list(reversed(records)), schema)
        self.assertEqual(first, second)
        self.assertEqual(first_report.as_dict(), second_report.as_dict())
        self.assertEqual([row[SOURCE_ID] for row in first], ["17", "29"])
        self.assertIsNone(first[0]["Email principal"])
        self.assertIsNone(first[0]["Nom de domaine / URL du lien"])
        self.assertEqual(first[0]["Produits"], None)

    def test_canonical_local_template_matches_sanitized_schema_fixture(self) -> None:
        canonical = ROOT / "templates" / "company.xlsx"
        if not canonical.exists():
            self.skipTest("local canonical SIMPLE export is not present")
        schema = load_simple_template_schema(canonical)
        self.assertEqual(schema.sheet_name, "Export")
        self.assertEqual(schema.headers, schema_headers())

    def _template(self) -> Path:
        directory = tempfile.mkdtemp()
        path = Path(directory) / "template.xlsx"
        make_template(path)
        self.addCleanup(lambda: shutil.rmtree(directory, ignore_errors=True))
        return path


if __name__ == "__main__":
    unittest.main()
