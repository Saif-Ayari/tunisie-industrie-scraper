from __future__ import annotations

import argparse
import json
from pathlib import Path
from unittest.mock import patch
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from crm.simple import load_simple_template_schema
from models import CompanyRecord
from scraper import export_both_outputs, run_export_state_mode
from state import CheckpointStore


ROOT = Path(__file__).resolve().parents[1]
SCHEMA_FIXTURE = ROOT / "tests" / "fixtures" / "simple_company_schema.json"


def make_template(path: Path) -> None:
    payload = json.loads(SCHEMA_FIXTURE.read_text(encoding="utf-8"))
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = payload["sheet"]
    worksheet.append(payload["headers"])
    workbook.save(path)


def args_for(directory: Path, state: Path | None = None) -> argparse.Namespace:
    template = directory / "simple-template.xlsx"
    make_template(template)
    return argparse.Namespace(
        output=directory / "raw.xlsx",
        simple_template=template,
        simple_output=directory / "simple.xlsx",
        export_state=state,
    )


def record(source_id: str) -> CompanyRecord:
    return CompanyRecord(
        company_name=f"Company {source_id}",
        short_name=f"C{source_id}",
        phone="(216) - 70 000 001 / 74 000 002",
        source_id=source_id,
        source_url=f"https://example.test/{source_id}",
        share_capital_dt="1\xa0500\xa0000",
        employees="10",
        scraped_at="2026-10-02T00:00:00Z",
    )


class Phase4Tests(unittest.TestCase):
    def test_shared_export_path_generates_raw_and_simple_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            args = args_for(directory_path)
            raw_summary, _, simple_summary = export_both_outputs([record("2"), record("1")], args)
            self.assertEqual(raw_summary["rows"], 2)
            self.assertEqual(simple_summary["rows"], 2)
            raw_workbook = load_workbook(args.output, data_only=False)
            simple_workbook = load_workbook(args.simple_output, data_only=False)
            raw = raw_workbook["Companies"]
            simple = simple_workbook["Export"]
            self.assertEqual([raw.cell(row, 18).value for row in range(2, 4)], ["1", "2"])
            simple_headers = tuple(cell.value for cell in simple[1])
            self.assertEqual(simple_headers, load_simple_template_schema(args.simple_template).headers)
            simple_id_column = simple_headers.index("ID source Tunisie Industrie") + 1
            self.assertEqual([simple.cell(row, simple_id_column).value for row in range(2, 4)], ["1", "2"])
            raw_workbook.close()
            simple_workbook.close()

    def test_export_state_uses_complete_accumulated_records_without_http(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            state_path = directory_path / "state"
            store = CheckpointStore(state_path)
            self.assertTrue(store.save_record(record("1")))
            resumed = CheckpointStore(state_path, resume=True)
            self.assertTrue(resumed.save_record(record("2")))
            args = args_for(directory_path, state_path)

            with patch("scraper.TunisieIndustrieClient", side_effect=AssertionError("HTTP must not be used")):
                self.assertEqual(run_export_state_mode(args), 0)

            raw_workbook = load_workbook(args.output, data_only=False)
            simple_workbook = load_workbook(args.simple_output, data_only=False)
            raw = raw_workbook["Companies"]
            simple = simple_workbook["Export"]
            self.assertEqual(raw.max_row - 1, 2)
            self.assertEqual(simple.max_row - 1, 2)
            raw_ids = {raw.cell(row, 18).value for row in range(2, raw.max_row + 1)}
            simple_headers = tuple(cell.value for cell in simple[1])
            simple_id_column = simple_headers.index("ID source Tunisie Industrie") + 1
            simple_ids = {simple.cell(row, simple_id_column).value for row in range(2, simple.max_row + 1)}
            self.assertEqual(raw_ids, {"1", "2"})
            self.assertEqual(simple_ids, raw_ids)
            self.assertEqual(len(simple_ids), 2)
            raw_workbook.close()
            simple_workbook.close()


if __name__ == "__main__":
    unittest.main()
