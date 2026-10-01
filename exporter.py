from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter

from models import CompanyRecord


HEADERS = (
    "Short Name",
    "Company Name",
    "Manager",
    "Activities",
    "Products",
    "Factory's Address",
    "Gouvernorate",
    "Delegation",
    "Phone Number Head Office/Factory",
    "Fax Number Head Office/Factory",
    "E-mail",
    "URL",
    "Market",
    "Foreign Participant Country",
    "Created",
    "Share Capital DT",
    "Employees",
    "Source ID",
    "Source URL",
    "Scraped At",
)

RECORD_ATTRIBUTES = (
    "short_name",
    "company_name",
    "manager",
    "activities",
    "products",
    "factory_address",
    "governorate",
    "delegation",
    "phone",
    "fax",
    "email",
    "url",
    "market",
    "foreign_participant_country",
    "created",
    "share_capital_dt",
    "employees",
    "source_id",
    "source_url",
    "scraped_at",
)


def safe_excel_text(value: object) -> str:
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def export_records(records: Iterable[CompanyRecord], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Companies"
    worksheet.append(list(HEADERS))
    for record in records:
        worksheet.append([safe_excel_text(getattr(record, attr)) for attr in RECORD_ATTRIBUTES])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{worksheet.max_row}"
    for index, header in enumerate(HEADERS, start=1):
        values = [worksheet.cell(row=row, column=index).value or "" for row in range(1, worksheet.max_row + 1)]
        width = min(max(len(str(value)) for value in values) + 2, 60)
        worksheet.column_dimensions[get_column_letter(index)].width = max(width, 14)
    workbook.save(output_path)


def verify_workbook(output_path: Path, expected_records: int) -> dict[str, object]:
    workbook = load_workbook(output_path, read_only=False, data_only=False)
    if "Companies" not in workbook.sheetnames:
        raise AssertionError("Expected Companies worksheet is missing")
    worksheet = workbook["Companies"]
    headers = tuple(cell.value for cell in worksheet[1])
    if headers != HEADERS:
        raise AssertionError(f"Unexpected headers: {headers!r}")
    if worksheet.max_row - 1 != expected_records:
        raise AssertionError(f"Expected {expected_records} data rows, got {worksheet.max_row - 1}")
    if worksheet.freeze_panes != "A2":
        raise AssertionError("Header row is not frozen")
    if worksheet.auto_filter.ref != f"A1:{get_column_letter(len(HEADERS))}{worksheet.max_row}":
        raise AssertionError("Autofilter range is missing or incorrect")
    source_url_index = HEADERS.index("Source URL") + 1
    scraped_at_index = HEADERS.index("Scraped At") + 1
    for row in range(2, worksheet.max_row + 1):
        if not worksheet.cell(row, source_url_index).value:
            raise AssertionError(f"Missing Source URL in row {row}")
        if not worksheet.cell(row, scraped_at_index).value:
            raise AssertionError(f"Missing Scraped At in row {row}")
    return {
        "worksheet": worksheet.title,
        "rows": worksheet.max_row - 1,
        "columns": worksheet.max_column,
        "freeze_panes": worksheet.freeze_panes,
        "autofilter": worksheet.auto_filter.ref,
    }

