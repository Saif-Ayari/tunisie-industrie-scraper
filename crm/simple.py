from __future__ import annotations

from dataclasses import dataclass, field
import re
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from models import CompanyRecord


# These names are copied from the inspected canonical SIMPLE export.  They are
# used to validate the supplied template; output column order always comes
# from that template rather than from this list.
DOMAIN_LABEL = "Nom de domaine / Libellé du lien"
DOMAIN_URL = "Nom de domaine / URL du lien"
DOMAIN_LINK = "Nom de domaine / Lien"
ADDRESS = "Adresse"
ADDRESS_2 = "Adresse 2"
CITY = "Ville"
STATE = "État"
COUNTRY = "Pays"
POSTAL_CODE = "Code postal"
LATITUDE = "Latitude"
LONGITUDE = "Longitude"
PHONE_NUMBER = "Numéro de téléphone"
PHONE_COUNTRY = "Code"
PHONE_CALLING_CODE = "Indicatif"
ADDITIONAL_PHONES = "Téléphone"
SOURCE_ID = "ID source Tunisie Industrie"
SOURCE_LABEL = "Source Tunisie Industrie / Libellé du lien"
SOURCE_URL = "Source Tunisie Industrie / URL du lien"
SOURCE_LINK = "Source Tunisie Industrie / Lien"
PRIMARY_EMAIL = "Email principal"
ADDITIONAL_EMAILS = "Courriels"

CRM_MANAGED_HEADERS = frozenset(
    {
        "ID de l'enregistrement",
        "Créé par",
        "Propriétaire du compte",
        "ID de Propriétaire du compte",
        "Date de création",
        "LinkedIn / Libellé du lien",
        "LinkedIn / URL du lien",
        "LinkedIn / Lien",
    }
)

REQUIRED_HEADERS = frozenset(
    {
        "Nom",
        DOMAIN_LABEL,
        DOMAIN_URL,
        DOMAIN_LINK,
        ADDRESS,
        ADDRESS_2,
        CITY,
        STATE,
        COUNTRY,
        POSTAL_CODE,
        LATITUDE,
        LONGITUDE,
        "Nom court",
        "Responsable",
        "Activités",
        "Produits",
        "Fax",
        "Marché",
        "Pays participant étranger",
        "Date de création entreprise",
        "Capital social DT",
        "Effectif",
        PHONE_NUMBER,
        PHONE_COUNTRY,
        PHONE_CALLING_CODE,
        ADDITIONAL_PHONES,
        SOURCE_ID,
        SOURCE_LABEL,
        SOURCE_URL,
        SOURCE_LINK,
        PRIMARY_EMAIL,
        ADDITIONAL_EMAILS,
    }
)


class SimpleExportError(ValueError):
    """Raised when a SIMPLE workbook cannot be safely generated."""


@dataclass(frozen=True)
class SimpleTemplateSchema:
    path: Path
    sheet_name: str
    headers: tuple[str, ...]

    @property
    def column_count(self) -> int:
        return len(self.headers)


@dataclass
class SimpleValidationReport:
    rows_transformed: int = 0
    warnings: list[str] = field(default_factory=list)
    blank_source_fields: dict[str, int] = field(default_factory=dict)
    duplicate_source_ids: list[str] = field(default_factory=list)

    def as_dict(self) -> dict[str, object]:
        return {
            "rows_transformed": self.rows_transformed,
            "warnings": list(self.warnings),
            "blank_source_fields": dict(sorted(self.blank_source_fields.items())),
            "duplicate_source_ids": list(self.duplicate_source_ids),
        }


@dataclass(frozen=True)
class _ParsedPhone:
    number: str
    country_code: str | None
    calling_code: str | None
    additional_serialization: str


@dataclass(frozen=True)
class _PhoneParts:
    primary: str | None
    country_code: str | None
    calling_code: str | None
    additional: tuple[str, ...]


def load_simple_template_schema(template_path: Path) -> SimpleTemplateSchema:
    """Read and validate the exact header contract from a SIMPLE export."""

    template_path = Path(template_path)
    if not template_path.is_file():
        raise SimpleExportError(f"SIMPLE template does not exist: {template_path}")

    try:
        workbook = load_workbook(template_path, read_only=True, data_only=False)
    except Exception as exc:  # noqa: BLE001 - expose a clear user-facing error.
        raise SimpleExportError(f"Could not read SIMPLE template {template_path}: {exc}") from exc

    try:
        if not workbook.worksheets:
            raise SimpleExportError(f"SIMPLE template has no worksheets: {template_path}")
        worksheet = workbook.active
        raw_headers = tuple(cell.value for cell in worksheet[1])
        if not raw_headers or any(not isinstance(value, str) or not value.strip() for value in raw_headers):
            raise SimpleExportError("SIMPLE template header row must contain only non-empty text headers")
        headers = tuple(value for value in raw_headers)
        duplicates = sorted({header for header in headers if headers.count(header) > 1})
        if duplicates:
            raise SimpleExportError(f"SIMPLE template has ambiguous duplicate headers: {duplicates!r}")
        missing = sorted(REQUIRED_HEADERS.difference(headers))
        if missing:
            raise SimpleExportError(f"SIMPLE template is missing required headers: {missing!r}")
        return SimpleTemplateSchema(path=template_path, sheet_name=worksheet.title, headers=headers)
    finally:
        workbook.close()


def _clean(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _split_compound_values(value: str | None) -> list[str]:
    text = _clean(value)
    if not text:
        return []
    return [part.strip() for part in re.split(r"\s*(?:/|;|\r?\n)\s*", text) if part.strip()]


def _is_tunisian_local_phone(value: str) -> bool:
    digits = re.sub(r"\D", "", value)
    return len(digits) == 8 and digits[0] in "23456789"


def _remove_tunisian_prefix(value: str) -> tuple[str, bool]:
    match = re.match(r"^\s*(?:\(\s*216\s*\)|\+216|216)\s*[-:]?\s*(.*)$", value)
    if match:
        return match.group(1).strip(), True
    return value.strip(), False


def _parse_phone_value(value: str) -> _ParsedPhone | None:
    number, had_tunisian_prefix = _remove_tunisian_prefix(value)
    if had_tunisian_prefix or _is_tunisian_local_phone(number):
        if not _is_tunisian_local_phone(number):
            return None
        return _ParsedPhone(
            number=number,
            country_code="TN",
            calling_code="+216",
            additional_serialization=f"+216{number} (TN)",
        )

    # Preserve an already-international number.  The CRM parser will resolve
    # its calling/country codes from the international representation.
    if re.fullmatch(r"\+\d[\d\s().-]*", value) and 7 <= len(re.sub(r"\D", "", value)) <= 15:
        return _ParsedPhone(
            number=value,
            country_code=None,
            calling_code=None,
            additional_serialization=value,
        )
    return None


def _phone_parts(value: str | None, report: SimpleValidationReport, source_id: str) -> _PhoneParts:
    values = _split_compound_values(value)
    if not values:
        return _PhoneParts(None, None, None, ())

    parsed_values: list[_ParsedPhone] = []
    for value_part in values:
        parsed = _parse_phone_value(value_part)
        if parsed is None:
            report.warnings.append(f"Source ID {source_id}: ignored malformed phone {value_part!r}")
            continue
        parsed_values.append(parsed)

    if not parsed_values:
        return _PhoneParts(None, None, None, ())

    first_source_value = values[0]
    first = _parse_phone_value(first_source_value)
    if first is None:
        return _PhoneParts(
            None,
            None,
            None,
            tuple(parsed.additional_serialization for parsed in parsed_values),
        )

    remaining = parsed_values[1:]
    return _PhoneParts(
        first.number,
        first.country_code,
        first.calling_code,
        tuple(parsed.additional_serialization for parsed in remaining),
    )


_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _email_parts(value: str | None, report: SimpleValidationReport, source_id: str) -> tuple[str | None, str | None]:
    values = _split_compound_values(value)
    valid: list[str] = []
    for email in values:
        if _EMAIL_RE.fullmatch(email):
            valid.append(email)
        else:
            report.warnings.append(f"Source ID {source_id}: ignored malformed email {email!r}")
    return (valid[0], "\n".join(valid[1:]) or None) if valid else (None, None)


_POSTAL_CODE_RE = re.compile(r"(?<!\d)(\d{4})(?!\d)")


def _postal_code(address: str | None) -> str | None:
    if not address:
        return None
    match = _POSTAL_CODE_RE.search(address)
    return match.group(1) if match else None


def _source_sort_key(record: CompanyRecord) -> tuple[int, object, str]:
    source_id = str(record.source_id)
    if source_id.isdigit():
        return (0, int(source_id), source_id)
    return (1, source_id, source_id)


def _put(row: dict[str, object], header: str, value: object) -> None:
    if value is not None:
        row[header] = value


def transform_records(
    records: Iterable[CompanyRecord],
    schema: SimpleTemplateSchema,
) -> tuple[list[dict[str, object]], SimpleValidationReport]:
    """Transform normalized source records into deterministic SIMPLE rows."""

    ordered = sorted(records, key=_source_sort_key)
    report = SimpleValidationReport()
    seen_ids: set[str] = set()
    rows: list[dict[str, object]] = []

    for record in ordered:
        source_id = _clean(record.source_id)
        if not source_id:
            raise SimpleExportError("Cannot export a company without Source ID")
        if source_id in seen_ids:
            report.duplicate_source_ids.append(source_id)
            raise SimpleExportError(f"Duplicate Source ID would be exported: {source_id}")
        seen_ids.add(source_id)

        row = {header: None for header in schema.headers}
        _put(row, "Nom", _clean(record.company_name))
        _put(row, "Nom court", _clean(record.short_name))
        _put(row, "Responsable", _clean(record.manager))
        _put(row, "Activités", _clean(record.activities))
        _put(row, "Produits", _clean(record.products))
        _put(row, "Fax", _clean(record.fax))
        _put(row, "Marché", _clean(record.market))
        _put(row, "Pays participant étranger", _clean(record.foreign_participant_country))
        _put(row, "Date de création entreprise", _clean(record.created))
        _put(row, "Capital social DT", _clean(record.share_capital_dt))
        _put(row, SOURCE_ID, source_id)

        address = _clean(record.factory_address)
        _put(row, ADDRESS, address)
        _put(row, CITY, _clean(record.delegation))
        _put(row, STATE, _clean(record.governorate))
        _put(row, POSTAL_CODE, _postal_code(address))

        phone = _phone_parts(record.phone, report, source_id)
        _put(row, PHONE_NUMBER, phone.primary)
        _put(row, PHONE_COUNTRY, phone.country_code)
        _put(row, PHONE_CALLING_CODE, phone.calling_code)
        _put(row, ADDITIONAL_PHONES, "\n".join(phone.additional) or None)

        primary_email, additional_emails = _email_parts(record.email, report, source_id)
        _put(row, PRIMARY_EMAIL, primary_email)
        _put(row, ADDITIONAL_EMAILS, additional_emails)

        company_url = _clean(record.url)
        _put(row, DOMAIN_URL, company_url)

        source_url = _clean(record.source_url)
        if source_url:
            _put(row, SOURCE_LABEL, "Tunisie Industrie")
            _put(row, SOURCE_URL, source_url)

        employees = _clean(record.employees)
        if employees:
            if re.fullmatch(r"\d+", employees):
                _put(row, "Effectif", int(employees))
            else:
                report.warnings.append(f"Source ID {source_id}: malformed Employees value {employees!r}; left blank")

        for field_name in ("company_name", "short_name", "manager", "activities", "products", "phone", "fax", "email"):
            if not _clean(getattr(record, field_name)):
                report.blank_source_fields[field_name] = report.blank_source_fields.get(field_name, 0) + 1

        rows.append(row)

    report.rows_transformed = len(rows)
    return rows, report


def _safe_cell_value(value: object) -> object:
    if isinstance(value, str) and value.startswith("="):
        return "'" + value
    return value


def export_simple_records(
    records: Iterable[CompanyRecord],
    template_path: Path,
    output_path: Path,
) -> SimpleValidationReport:
    schema = load_simple_template_schema(template_path)
    rows, report = transform_records(records, schema)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = schema.sheet_name
    worksheet.append(list(schema.headers))
    for row in rows:
        worksheet.append([_safe_cell_value(row.get(header)) for header in schema.headers])

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(schema.column_count)}{worksheet.max_row}"
    for index in range(1, schema.column_count + 1):
        values = [worksheet.cell(row=row, column=index).value or "" for row in range(1, worksheet.max_row + 1)]
        width = min(max(len(str(value)) for value in values) + 2, 60)
        worksheet.column_dimensions[get_column_letter(index)].width = max(width, 14)
        for row_number in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_number, column=index)
            if isinstance(cell.value, str):
                cell.number_format = "@"
    workbook.save(output_path)
    verify_simple_workbook(output_path, schema, len(rows))
    return report


def verify_simple_workbook(
    output_path: Path,
    schema: SimpleTemplateSchema,
    expected_records: int,
) -> dict[str, object]:
    workbook = load_workbook(output_path, read_only=False, data_only=False)
    if schema.sheet_name not in workbook.sheetnames:
        raise AssertionError(f"Expected worksheet {schema.sheet_name!r} is missing")
    worksheet = workbook[schema.sheet_name]
    headers = tuple(cell.value for cell in worksheet[1])
    if headers != schema.headers:
        raise AssertionError(f"Unexpected SIMPLE headers: {headers!r}")
    if worksheet.max_row - 1 != expected_records:
        raise AssertionError(f"Expected {expected_records} SIMPLE rows, got {worksheet.max_row - 1}")
    if worksheet.max_column != schema.column_count:
        raise AssertionError(f"Expected {schema.column_count} SIMPLE columns, got {worksheet.max_column}")
    if worksheet.freeze_panes != "A2":
        raise AssertionError("SIMPLE header row is not frozen")
    expected_filter = f"A1:{get_column_letter(schema.column_count)}{worksheet.max_row}"
    if worksheet.auto_filter.ref != expected_filter:
        raise AssertionError(f"Unexpected SIMPLE autofilter: {worksheet.auto_filter.ref!r}")

    source_id_column = schema.headers.index(SOURCE_ID) + 1
    source_ids: list[str] = []
    for row_number in range(2, worksheet.max_row + 1):
        value = worksheet.cell(row_number, source_id_column).value
        if not value:
            raise AssertionError(f"Missing Source ID in SIMPLE row {row_number}")
        source_ids.append(str(value))
    if len(source_ids) != len(set(source_ids)):
        raise AssertionError("Duplicate Source IDs are present in the SIMPLE workbook")

    for header in CRM_MANAGED_HEADERS:
        if header not in schema.headers:
            continue
        column = schema.headers.index(header) + 1
        for row_number in range(2, worksheet.max_row + 1):
            if worksheet.cell(row_number, column).value not in (None, ""):
                raise AssertionError(f"CRM-managed header {header!r} is populated in row {row_number}")

    summary = {
        "worksheet": worksheet.title,
        "rows": worksheet.max_row - 1,
        "columns": worksheet.max_column,
        "freeze_panes": worksheet.freeze_panes,
        "autofilter": worksheet.auto_filter.ref,
        "source_ids": source_ids,
    }
    workbook.close()
    return summary
